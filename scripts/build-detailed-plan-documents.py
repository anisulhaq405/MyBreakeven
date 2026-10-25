"""Build the detailed and five year editable plans from shared resource data."""
import json, subprocess
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parents[1]
raw=subprocess.check_output(['node','--input-type=module','-e',"import{planCase,generalSections,yearMilestones,annualPlanRows}from'./src/detailedPlanData.js';console.log(JSON.stringify({planCase,generalSections,yearMilestones,rows:annualPlanRows()}))"],cwd=root)
data=json.loads(raw);sections=data['generalSections'];rows=data['rows']
def money(v):return ('−$'+format(-v,',.0f')) if v<0 else '$'+format(v,',.0f')
def document(title):
 d=Document();s=d.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.top_margin=s.bottom_margin=Inches(.65);s.left_margin=s.right_margin=Inches(.7)
 for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
  d.styles[name].font.name='Arial';d.styles[name].font.color.rgb=RGBColor(0,0,0)
 d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.space_after=Pt(6);d.styles['Normal'].paragraph_format.line_spacing=1.05
 d.styles['Title'].font.size=Pt(22)
 for name in ['Heading 1','Heading 2']:
  d.styles[name].font.size=Pt(12);d.styles[name].paragraph_format.space_before=Pt(9);d.styles[name].paragraph_format.space_after=Pt(3)
 for style in d.styles:
  for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
 d.core_properties.title=title;d.core_properties.author='MyBreakeven'
 foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER;foot.add_run('MyBreakeven  |  ')
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
 for run in foot.runs:run.font.size=Pt(8)
 return d
def table(d,headers,body,widths=None):
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False
 if widths:
  for col,width in zip(t.columns,widths):col.width=Inches(width)
 for i,v in enumerate(headers):t.rows[0].cells[i].text=str(v)
 pr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');pr.append(repeat)
 for row in body:
  for cell,v in zip(t.add_row().cells,row):cell.text=str(v)
 for ri,row in enumerate(t.rows):
  for ci,cell in enumerate(row.cells):
   cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;tc=cell._tc.get_or_add_tcPr()
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'17314F' if ri==0 else ('EDF2F7' if ri%2 else 'FFFFFF'));tc.append(shade)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    border=OxmlElement('w:'+edge);border.set(qn('w:val'),'single');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'D9D9D9');borders.append(border)
   tc.append(borders);marg=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');marg.append(e)
   tc.append(marg)
   if widths:cell.width=Inches(widths[ci])
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
    if ci>0 and len(headers)>2 and headers[0]=='Measure':p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    for run in p.runs:run.font.size=Pt(9.5);run.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0);run.bold=ri==0
 d.add_paragraph('')
 return t
def start(d,title,sub):d.add_heading(title,0);d.add_paragraph(sub)
def page(d,title,sub):d.add_page_break();start(d,title,sub)
def section(d,title,text):d.add_heading(title,1);d.add_paragraph(text)

d=document('Business Plan Template');start(d,'Business Plan Template','Blank editable outline and completed fictional example  |  October 25 2026')
d.add_paragraph('Use this document to write the business decision and the evidence behind it. Replace every prompt and fictional claim. The example is a small service business, not a universal price or funding benchmark.')
d.add_paragraph('Business name  ____________________   Owner  ____________________')
d.add_paragraph('Planning period  ____________________   Next review  ____________________')
for title,prompt,_ in sections[:5]:section(d,title,prompt+' [Write your response and attach evidence.]')
page(d,'Business Plan Outline','Continue the blank plan and record unresolved evidence')
for title,prompt,_ in sections[5:]:section(d,title,prompt+' [Write your response and attach evidence.]')
section(d,'Before sharing the plan','Check the audience and decision requested. Match the narrative to the financial schedules. Confirm the opening cash cutoff and record what remains unverified. A lender or investor may request statements, historical records and a different format.')
page(d,'Harbor Desk Services','Completed fictional business plan for calendar year 2027')
for title,_,sample in sections[:5]:section(d,title,sample)
page(d,'Harbor Desk Services Operating Plan','Completed fictional example continued')
for title,_,sample in sections[5:]:section(d,title,sample)
page(d,'Business Plan Financial Example','Fictional USD assumptions for 2027 with a transparent cash boundary')
table(d,['Measure','Checked example'],[
 ['Revenue','120 packages × $500 = $60,000'],['Contribution','$60,000 − $18,000 direct costs = $42,000'],['Planning result','$42,000 − $12,000 fixed cash − $24,000 owner allowance = $6,000'],['Whole packages for commitments','Round up $36,000 ÷ $350 = 103 packages'],['Capacity','1,200 productive team hours ÷ 8 = 150 packages'],['Cash collected','$60,000 × 90% = $54,000'],['Year end cash','$15,000 + $54,000 − $54,000 operating and owner cash − $3,000 equipment = $12,000'],['Closing receivables','$60,000 × 10% = $6,000']], [2.1,5])
section(d,'Write the executive summary after checking the plan','Summarize the specific buyer and offer, the evidence collected, the operating and cash requirement, the funding source and the next approval. In this example the proposed $15,000 opening cash is not an approved financing offer. The owner must demonstrate that it is available and test the payment dates.')
section(d,'Attach the schedules required by the reader','Use the companion five year resource for the annual planning assumptions. Build monthly or quarterly first year projections and attach the income statements, balance sheets, cash flow statements and supporting schedules requested by the recipient. This narrative and simplified example do not supply an integrated financial statement model.')
section(d,'Avoid duplicate costs','Owner cash support is modeled once as a separate allowance. Direct costs already include the stated support, transaction and acquisition charges. Equipment cash is separate from operating costs. Tax, debt service, depreciation and other unentered amounts are not silently assumed to be covered.')
d.add_paragraph('Resource  https://mybreakeven.com/resources/business-plan-template/')
d.add_paragraph('Outline reference  U.S. Small Business Administration')
d.add_paragraph('https://www.sba.gov/counseling/plan-your-business/')
d.save(root/'public/downloads/business-plan-template.docx')

d=document('Five Year Business Plan Template');start(d,'Five Year Business Plan Template','Blank strategy outline and completed fictional case  |  October 25 2026')
d.add_paragraph('Connect a five year direction to annual targets, delivery resources, cash and evidence. The Excel worksheet models annual planning relationships. It is not an integrated balance sheet, income statement and cash flow model or a funding approval.')
d.add_paragraph('Business  ____________________   Period  ____________________')
for title,text in [
 ('Starting position','Describe current products, customers, team, finances and verified constraints. [Replace with evidence.]'),
 ('Five year direction','Name the intended customer and offer, the purpose of growth and the conditions that would change the direction. [Write your decision.]'),
 ('Annual objectives','Choose a measurable objective for each year and the evidence needed before expanding. [Complete the annual table.]'),
 ('Delivery resources','Reconcile planned volume, productive team hours, costed support and practical capacity. [Attach the operating schedule.]'),
 ('Financial assumptions','List annual volume, price, direct costs, fixed commitments, owner cash and collection timing. [Replace Excel blue inputs.]'),
 ('Funding and cash','State opening cash, capital purchases, other payments, new funding and the chosen floor. [Attach actual payment dates.]'),
 ('Review and response','Assign owners and review dates. State what spending or expansion stops if a gate fails. [Complete the evidence register.]')]:section(d,title,text)
page(d,'Annual Objectives Template','Replace prompts with measurable outcomes and assigned reviews')
table(d,['Year','Objective and measurable target','Owner','Next review'],[[f'Year {n}','[Target and evidence]','[Assign]','[Date]'] for n in range(1,6)],[.7,3.7,1.3,1.4])
section(d,'Evidence for each assumption','Retain dated customer records, accepted scopes, supplier quotes, complete delivery hour logs and collection terms. A later year target needs its own evidence and resources. Increasing a volume cell does not establish demand or hire another worker.')
section(d,'Approval record','Decision  ____________________   Approved by  ____________________')
d.add_paragraph('Evidence reviewed  __________________________________________________')
d.add_paragraph('Unresolved condition  _________________________________________________')
d.add_paragraph('Action and owner  ___________________________________________________')
section(d,'First year detail','Keep a monthly or quarterly first year schedule alongside the annual plan. Use the dated cash forecast for supplier, payroll and settlement dates. Annual positive cash can conceal a shortage before a customer pays.')
page(d,'Harbor Desk Services Five Year Plan','Completed fictional case for 2027 to 2031  |  USD')
d.add_paragraph('The owner will expand a defined document service only after paid tests validate demand, complete delivery hours and collection terms. Annual targets are entered separately rather than generated from an assumed market growth rate. Direct delivery support and the owner allowance must fund the team hours shown.')
table(d,['Measure']+[str(r['year']) for r in rows],[[label]+[formatter(r[key]) for r in rows] for label,key,formatter in [('Packages','jobs',str),('Price','price',money),('Revenue','revenue',money),('Contribution','contribution',money),('After owner allowance','result',money),('Package capacity','capacity',str)]],[2.1,1,1,1,1,1])
for year,objective,target,evidence,action in data['yearMilestones']:
 section(d,year+' '+objective,target+'. '+evidence+'. '+action+'.')
page(d,'Five Year Cash and Review','The annual cash bridge uses the same inputs as the Excel worksheet')
table(d,['Measure']+[str(r['year']) for r in rows],[[label]+[money(r[key]) for r in rows] for label,key in [('Opening cash','openingCash'),('Customer cash','receipts'),('Operating and owner cash','payments'),('Equipment cash','capex'),('Closing cash','closingCash'),('Uncollected sales','closingAR')]],[2.1,1,1,1,1,1])
section(d,'Collection and cost assumptions','The case starts with $15,000 cash and no opening receivables. Each year collects 90% of current sales plus all prior year receivables. All entered direct costs, fixed commitments and owner allowance are paid in that year. No bad debt or delay beyond the next year is modeled. Additional funding and other cash payments start at zero, not an assertion that tax or debt obligations do not exist.')
section(d,'A first year downside test','At 90 packages with unchanged price, costs and owner allowance, the planning result is −$4,500 and year end cash is $3,000 after equipment. This is below the illustrative $5,000 floor. Delay discretionary spending or revise scope and funding before taking on commitments. Test the actual dates before using even the base case.')
section(d,'Update the workbook and the written plan together','In Annual plan, edit blue inputs and retain the formulas. The following year rolls opening cash and receivables from the previous year. In Milestones, replace example goals and add an owner and next review date. Check each gate against evidence before approving expansion. Reconcile monthly actuals and formally review the remaining annual targets each quarter.')
d.add_paragraph('Resource  https://mybreakeven.com/resources/five-year-business-plan-template/')
d.save(root/'public/downloads/five-year-business-plan-template.docx')
print('Two editable business plans written')
