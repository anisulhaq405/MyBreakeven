"""Build an independent 12-month restaurant budget workbook (not the monthly download).
Runs in CI with openpyxl. Formulas calculate when opened in Excel/Sheets.
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import LineChart, Reference
from openpyxl.utils import get_column_letter

OUTPUT=Path("public/downloads/restaurant-12-month-budget-template.xlsx")
wb=Workbook()
read=wb.active; read.title="Instructions"
read.append(["MYBREAKEVEN — 12-MONTH RESTAURANT BUDGET", ""])
for line in [
("Purpose","Monthly budget, monthly actuals and year-to-date variance; USD illustration only"),
("Input","Replace sample budget assumptions in yellow cells; enter actuals after each month closes"),
("Important","Blank actual months do not count as zero actual profits"),
("Completion","Enter all actual income and cost rows, including actual zeros, before interpreting results"),
("Profit","Sales less variable costs and committed expenses; not cash flow or after-tax earnings"),
("Break-even","Committed expenses / contribution margin, assuming constant variable-cost ratio"),
("Comparison","Annual budget spans 12 months; actual and variance sum entered months only"),
("Caution","Payroll excludes employer costs when listed separately; do not double count owner compensation"),
("Example","Illustrative amounts, not industry benchmarks"),
("Original page","https://mybreakeven.com/blogs/restaurant-budgeting/"),
]: read.append(list(line))
read.column_dimensions["A"].width=30;read.column_dimensions["B"].width=104
read["A1"].font=Font(bold=True,color="FFFFFF",size=13)
read["A1"].fill=PatternFill("solid",fgColor="163754")
for row in read.iter_rows(min_row=2,max_col=2): row[0].font=Font(bold=True,color="163754");row[1].alignment=Alignment(wrap_text=True)
plan=wb.create_sheet("12-Month Budget")
plan["A1"]="RESTAURANT — MONTHLY BUDGET / ACTUAL (USD)"
plan["A1"].font=Font(bold=True,color="FFFFFF",size=15);plan["A1"].fill=PatternFill("solid",fgColor="163754")
months=["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]
for j,label in enumerate(["Line item"]+months+["Annual / entered months"],1):
 cell=plan.cell(3,j,label);cell.fill=PatternFill("solid",fgColor="163754");cell.font=Font(bold=True,color="FFFFFF")
rows={4:"BUDGET",5:"Food sales",6:"Beverage sales",7:"Total sales",8:"Food and beverage COGS",9:"Packaging/delivery/fees (variable)",10:"Contribution",11:"Guaranteed payroll",12:"Employer payroll costs",13:"Owner compensation",14:"Occupancy",15:"Utilities",16:"Marketing",17:"Other committed expenses",18:"Committed expenses",19:"Budget operating profit",20:"Contribution margin",21:"Break-even sales",23:"ACTUAL",24:"Food sales",25:"Beverage sales",26:"Total sales",27:"Food and beverage COGS",28:"Packaging/delivery/fees (variable)",29:"Contribution",30:"Guaranteed payroll",31:"Employer payroll costs",32:"Owner compensation",33:"Occupancy",34:"Utilities",35:"Marketing",36:"Other committed expenses",37:"Committed expenses",38:"Actual operating profit",39:"Contribution margin",40:"Break-even sales",42:"VARIANCE: ACTUAL MINUS BUDGET",43:"Sales",44:"Variable costs",45:"Committed expenses",46:"Operating profit"}
for r,label in rows.items():
 plan.cell(r,1,label)
 if r in (4,23,42):
  for col in range(1,15):
   cell=plan.cell(r,col);cell.fill=PatternFill("solid",fgColor="235E80");cell.font=Font(color="FFFFFF",bold=True)
for i in range(12):
 col=get_column_letter(i+2);factor=[.90,.91,.98,1,1.05,1.07,1.12,1.08,1.02,1.01,1.04,1.15][i]
 sample={5:round(45000*factor),6:round(15000*factor),8:round(18000*factor),9:round(1800*factor),11:15000,12:1500,13:4000,14:6000,15:1200,16:800,17:1700}
 for r,val in sample.items(): plan[f"{col}{r}"]=val
 formulas={7:f"=SUM({col}5:{col}6)",10:f"={col}7-SUM({col}8:{col}9)",18:f"=SUM({col}11:{col}17)",19:f"={col}10-{col}18",20:f'=IF({col}7=0,"",{col}10/{col}7)',21:f'=IF({col}7=0,"No sales",IF({col}10<=0,"No positive contribution",{col}18/({col}10/{col}7)))',
 26:f'=IF(COUNT({col}24:{col}25)=0,"",SUM({col}24:{col}25))',
 29:f'=IF({col}26="","",{col}26-SUM({col}27:{col}28))',
 37:f'=IF(COUNT({col}30:{col}36)=0,"",SUM({col}30:{col}36))',
 38:f'=IF(OR({col}26="",{col}37=""),"",{col}29-{col}37)',
 39:f'=IF(OR({col}26="",{col}26=0),"",{col}29/{col}26)',
 40:f'=IF({col}26="","",IF({col}26=0,"No sales",IF({col}29<=0,"No positive contribution",{col}37/({col}29/{col}26))))',
 43:f'=IF({col}38="","",{col}26-{col}7)',
 44:f'=IF({col}38="","",SUM({col}27:{col}28)-SUM({col}8:{col}9))',
 45:f'=IF({col}38="","",{col}37-{col}18)',
 46:f'=IF({col}38="","",{col}38-{col}19)'}
 for r,formula in formulas.items():plan[f"{col}{r}"]=formula
 for r in list(sample)+[24,25,27,28,30,31,32,33,34,35,36]:
  cell=plan[f"{col}{r}"];cell.fill=PatternFill("solid",fgColor="FFF1BD");cell.font=Font(color="1646BB")
for r in [5,6,7,8,9,10,11,12,13,14,15,16,17,18,19,24,25,26,27,28,29,30,31,32,33,34,35,36,37,38,43,44,45,46]:
 plan[f"N{r}"]=f"=SUM(B{r}:M{r})"
for r,num,den in [(20,10,7),(39,29,26)]:
 plan[f"N{r}"]=f'=IF(N{den}=0,"",N{num}/N{den})'
for r,sales,contrib,fixed in [(21,7,10,18),(40,26,29,37)]:
 plan[f"N{r}"]=f'=IF(N{sales}=0,"No sales",IF(N{contrib}<=0,"No positive contribution",N{fixed}/(N{contrib}/N{sales})))'
for r in [7,10,18,19,26,29,37,38,46]:
 for cell in plan[r][:14]:cell.fill=PatternFill("solid",fgColor="EAF3FA")
for row in plan.iter_rows(min_row=5,max_row=46,min_col=2,max_col=14):
 for cell in row:cell.number_format='$#,##0;[Red]($#,##0);-'
for r in (20,39):
 for cell in plan[r][1:14]:cell.number_format="0.0%"
plan.column_dimensions["A"].width=43
for j in range(2,15):plan.column_dimensions[get_column_letter(j)].width=16
plan.freeze_panes="B4"
dash=wb.create_sheet("Dashboard")
dash.append(["RESTAURANT BUDGET — ANNUAL SUMMARY","USD / ratio"])
for label,src in [("Budget sales",7),("Actual sales (entered months)",26),("Budget operating profit",19),("Actual operating profit (entered months)",38),("Operating profit variance (entered months)",46),("Budget contribution margin",20),("Actual contribution margin",39)]:
 dash.append([label,f"='12-Month Budget'!N{src}"])
for cell in dash[1][:2]:cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="163754")
dash["A11"]="Month";dash["B11"]="Budget profit";dash["C11"]="Actual profit";dash["D11"]="Variance"
for i,month in enumerate(months):
 row=i+12;col=get_column_letter(i+2)
 dash.cell(row,1,month)
 for n,source in [(2,19),(3,38),(4,46)]:dash.cell(row,n,f"='12-Month Budget'!{col}{source}")
chart=LineChart();chart.title="Monthly operating profit: budget vs actual"
chart.add_data(Reference(dash,min_col=2,max_col=3,min_row=11,max_row=23),titles_from_data=True)
chart.set_categories(Reference(dash,min_col=1,min_row=12,max_row=23))
dash.add_chart(chart,"F4")
dash.column_dimensions["A"].width=44
for col in "BCD":dash.column_dimensions[col].width=22
OUTPUT.parent.mkdir(parents=True,exist_ok=True)
wb.save(OUTPUT)
print(f"Built {OUTPUT}")
