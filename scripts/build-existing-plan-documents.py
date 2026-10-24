"""Generate the three editable narrative plans from their shared page content."""
import json, subprocess
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parents[1]
raw=subprocess.check_output(['node','--input-type=module','-e',"import {existingPlanNarratives} from './src/existingPlanNarratives.js';console.log(JSON.stringify(existingPlanNarratives))"],cwd=root)
plans=json.loads(raw)
for key,p in plans.items():
 d=Document(); sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
 sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
 normal=d.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10.5);normal.paragraph_format.space_after=Pt(6)
 normal.paragraph_format.line_spacing=1.05
 if key=='freelance':
  normal.font.size=Pt(10);normal.paragraph_format.space_after=Pt(4)
 for style in ['Heading 1','Heading 2']:
  d.styles[style].font.name='Arial';d.styles[style].font.color.rgb=RGBColor.from_string('174C59')
  d.styles[style].font.size=Pt(12);d.styles[style].paragraph_format.space_before=Pt(9);d.styles[style].paragraph_format.space_after=Pt(3)
 d.styles['Title'].font.name='Arial';d.styles['Title'].font.size=Pt(23);d.styles['Title'].font.color.rgb=RGBColor.from_string('15374F')
 for st in d.styles:
  for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
 d.core_properties.title=p['title'];d.core_properties.author='MyBreakeven';d.core_properties.subject=p['subtitle']
 d.add_heading(p['title'],0);d.add_paragraph('Blank editable outline | MyBreakeven | October 24, 2026')
 d.add_paragraph('Business name: ____________________   Owner: ____________________')
 d.add_paragraph('Plan period: ____________________   Next review: ____________________')
 d.add_heading('Executive summary',1);d.add_paragraph('Summarize the buyer, offer, evidence, funding need and next milestone. Complete this after checking the financial schedules.')
 for title,prompt,_ in p['sections']:
  d.add_heading(title,1);d.add_paragraph(prompt+' [Replace with your plan.]')
 d.add_page_break();d.add_heading(p['name'],0);d.add_paragraph('Completed fictional example | '+p['subtitle'])
 d.add_paragraph('Executive summary: '+p['sections'][0][2]+' '+p['sections'][5 if key=='freelance' else 6][2])
 for title,_,sample in p['sections']:
  d.add_heading(title,1);d.add_paragraph(sample)
 d.add_page_break();d.add_heading('Financial example and evidence',0)
 d.add_paragraph('USD example for one modeled month. Use the companion Excel workbook to replace inputs and inspect the operating calculations. These are fictional inputs, not industry averages or customer results.')
 t=d.add_table(rows=1,cols=2);t.style='Light Shading Accent 1';t.rows[0].cells[0].text='Measure';t.rows[0].cells[1].text='Calculation'
 for label,value in p['checks']:
  cells=t.add_row().cells;cells[0].text=label;cells[1].text=value
 trPr=t.rows[0]._tr.get_or_add_trPr(); repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
 d.add_paragraph(p['explanation']);d.add_heading('Cash timing',1);d.add_paragraph(p['cash'])
 d.add_heading('Evidence to collect',1)
 for x in p['evidence']:d.add_paragraph(x,style='List Bullet')
 d.add_paragraph('Record each next action with an owner, due date, evidence and a pass condition. Replace assumptions after paid tests and reconcile completed work, invoices and bank receipts each month.')
 d.add_paragraph('Scope: results exclude unentered tax, financing, depreciation and owner costs. Monthly cash does not test payment dates within a month. No external accountant review or funding approval is implied.')
 d.add_paragraph('Resource: https://mybreakeven.com/blogs/'+p['slug']+'/')
 d.save(root/'public'/'downloads'/p['document'])
 print(key,'written')
